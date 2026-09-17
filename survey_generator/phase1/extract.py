"""
Phase 1 — file extraction.

Uploaded files are converted into bounded, labelled plain text before the LLM
sees them. Each extractor returns metadata for the extraction summary: file type,
pages/slides/sheets processed, character counts, truncation, weak extraction,
and warnings/errors. Legacy binary Office files (.doc, .ppt, .xls) remain
unsupported; users should convert them to .docx, .pptx, or .xlsx first.
"""
from __future__ import annotations

import csv as _csv
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from typing import Iterable, List
from xml.etree import ElementTree as ET

from .models import SourceDoc
from survey_generator.runtime import excel_safe_row, sanitize_text, shorten_text

PER_FILE_CHAR_LIMIT = int(os.environ.get("PHASE1_PER_FILE_CHAR_LIMIT", "20000"))
PREVIEW_CHAR_LIMIT = int(os.environ.get("PHASE1_PREVIEW_CHAR_LIMIT", "5000"))
WEAK_EXTRACTION_CHAR_THRESHOLD = int(os.environ.get("PHASE1_WEAK_EXTRACTION_CHAR_THRESHOLD", "500"))
ENABLE_OCR = os.environ.get("PHASE1_ENABLE_OCR", "0").strip().lower() in {"1", "true", "yes", "on"}
EXCEL_MAX_ROWS_PER_SHEET = int(os.environ.get("PHASE1_EXCEL_MAX_ROWS_PER_SHEET", "120"))
CSV_MAX_ROWS = int(os.environ.get("PHASE1_CSV_MAX_ROWS", "250"))


@dataclass
class ExtractedContent:
    kind: str
    text: str = ""
    item_count: int = 0
    item_unit: str = "items"
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ocr_attempted: bool = False
    ocr_used: bool = False


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower().lstrip(".")


def _clean_text(text: str) -> str:
    # Normalize text as early as possible. PPTX speaker notes and copied slide
    # titles can contain invisible control characters such as vertical-tab
    # (\x0b), which later crash XLSX exports if not removed.
    text = sanitize_text(text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _source_aware_truncate(text: str, limit: int = PER_FILE_CHAR_LIMIT) -> tuple[str, bool, str]:
    """Preserve beginning, middle, and end rather than chopping off the tail."""
    text = _clean_text(text)
    if len(text) <= limit:
        return text, False, ""
    marker1 = "\n\n[... SOURCE-AWARE TRUNCATION: middle excerpt follows; intervening text omitted ...]\n\n"
    marker2 = "\n\n[... SOURCE-AWARE TRUNCATION: final excerpt follows ...]\n\n"
    marker_budget = len(marker1) + len(marker2)
    available = max(limit - marker_budget, max(120, limit // 2))
    start_len = max(int(available * 0.45), 1)
    mid_len = max(int(available * 0.25), 1)
    end_len = max(available - start_len - mid_len, 1)
    # If markers plus minimum excerpts would still exceed the limit, shrink the
    # excerpts proportionally rather than slicing off the tail after assembly.
    while start_len + mid_len + end_len + marker_budget > limit and end_len > 1:
        end_len -= 1
    while start_len + mid_len + end_len + marker_budget > limit and mid_len > 1:
        mid_len -= 1
    while start_len + mid_len + end_len + marker_budget > limit and start_len > 1:
        start_len -= 1
    midpoint = max((len(text) - mid_len) // 2, start_len)
    truncated = f"{text[:start_len].rstrip()}{marker1}{text[midpoint:midpoint + mid_len].strip()}{marker2}{text[-end_len:].lstrip()}"
    warning = f"Truncated source text from {len(text):,} to {len(truncated):,} characters using beginning/middle/end excerpts."
    return truncated, True, warning


def _xml_text_nodes(xml_bytes: bytes) -> list[str]:
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return []
    out: list[str] = []
    for node in root.iter():
        if node.tag.endswith("}t") or node.tag == "t":
            cleaned = _clean_text(node.text or "")
            if cleaned:
                out.append(cleaned)
    return out


def _ocr_pdf_best_effort(data: bytes) -> ExtractedContent:
    if not ENABLE_OCR:
        return ExtractedContent(
            kind="pdf",
            text="",
            warnings=["OCR is disabled. Set PHASE1_ENABLE_OCR=1 and install OCR dependencies to process scanned PDFs."],
            ocr_attempted=False,
        )
    try:
        from pdf2image import convert_from_bytes  # type: ignore
        import pytesseract  # type: ignore
    except Exception as exc:  # noqa: BLE001
        return ExtractedContent(kind="pdf", text="", warnings=[f"OCR was enabled but OCR dependencies are unavailable: {exc}"], ocr_attempted=True)
    try:
        pages = convert_from_bytes(data)
        parts: list[str] = []
        for i, image in enumerate(pages, start=1):
            text = pytesseract.image_to_string(image) or ""
            if text.strip():
                parts.append(f"--- OCR Page {i} ---\n{text.strip()}")
        return ExtractedContent(
            kind="pdf",
            text="\n\n".join(parts),
            item_count=len(pages),
            item_unit="pages OCRed",
            warnings=["OCR text was generated from page images and may contain recognition errors."],
            ocr_attempted=True,
            ocr_used=bool(parts),
        )
    except Exception as exc:  # noqa: BLE001
        return ExtractedContent(kind="pdf", text="", warnings=[f"OCR failed: {exc}"], ocr_attempted=True)


# ---- individual extractors --------------------------------------------------
def extract_pdf(data: bytes) -> ExtractedContent:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    page_count = len(reader.pages)
    parts: list[str] = []
    zero_text_pages = 0
    warnings: list[str] = []
    errors: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            page_text = _clean_text(page.extract_text() or "")
            if page_text:
                parts.append(f"--- Page {i} ---\n{page_text}")
            else:
                zero_text_pages += 1
        except Exception as exc:  # noqa: BLE001
            zero_text_pages += 1
            errors.append(f"Page {i}: could not extract text ({exc})")
    text = "\n\n".join(parts)
    char_count = len(_clean_text(text))
    ocr_attempted = False
    ocr_used = False
    if page_count and zero_text_pages:
        warnings.append(f"{zero_text_pages} of {page_count} PDF page(s) had no extractable text.")
    if page_count and (char_count == 0 or char_count / max(page_count, 1) < 50):
        warnings.append("This PDF appears scanned or image-heavy; text extraction may be incomplete.")
        ocr = _ocr_pdf_best_effort(data)
        ocr_attempted = ocr.ocr_attempted
        ocr_used = ocr.ocr_used
        if ocr.text.strip():
            text = (text + "\n\n" + ocr.text).strip()
        warnings.extend(ocr.warnings)
        errors.extend(ocr.errors)
    return ExtractedContent(kind="pdf", text=text, item_count=page_count, item_unit="pages", warnings=warnings, errors=errors, ocr_attempted=ocr_attempted, ocr_used=ocr_used)


def extract_docx(data: bytes) -> ExtractedContent:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    table_count = 0
    for tbl in doc.tables:
        table_count += 1
        parts.append(f"--- Table {table_count} ---")
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    warnings = [] if parts else ["DOCX contained no extractable paragraph or table text."]
    return ExtractedContent(kind="docx", text="\n".join(parts), item_count=len(doc.paragraphs) + table_count, item_unit="paragraphs/tables", warnings=warnings)


def _pptx_notes_text(data: bytes) -> list[str]:
    notes: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = sorted(
                [n for n in zf.namelist() if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")],
                key=lambda n: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", n)],
            )
            for name in names:
                bits = [b for b in _xml_text_nodes(zf.read(name)) if b.lower() not in {"click to edit notes", "notes"}]
                if bits:
                    notes.append("\n".join(bits))
    except Exception:  # noqa: BLE001
        return []
    return notes


def extract_pptx(data: bytes) -> ExtractedContent:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    notes = _pptx_notes_text(data)
    parts: list[str] = []
    warnings: list[str] = []
    image_only_slides = 0
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        try:
            title = getattr(slide.shapes, "title", None)
            if title is not None and getattr(title, "has_text_frame", False):
                title_text = _clean_text(title.text)
                if title_text:
                    slide_lines.append(f"Title: {title_text}")
            for shape in slide.shapes:
                try:
                    if getattr(shape, "has_text_frame", False):
                        text = _clean_text("\n".join(p.text for p in shape.text_frame.paragraphs if p.text and p.text.strip()))
                        if text and f"Title: {text}" not in slide_lines:
                            slide_lines.append(text)
                    if getattr(shape, "has_table", False):
                        slide_lines.append("Table:")
                        for row in shape.table.rows:
                            cells = [_clean_text(c.text) for c in row.cells if c.text and _clean_text(c.text)]
                            if cells:
                                slide_lines.append(" | ".join(cells))
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"Slide {i}: one shape could not be read ({shorten_text(exc, 180)}).")
            if i <= len(notes) and notes[i - 1].strip():
                slide_lines.append("Speaker notes:\n" + _clean_text(notes[i - 1]))
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"Slide {i}: could not be fully read ({shorten_text(exc, 180)}).")
        if not slide_lines:
            image_only_slides += 1
        parts.append(f"--- Slide {i} ---")
        parts.extend(slide_lines or ["(no extractable slide text)"])
    if image_only_slides:
        warnings.append(f"{image_only_slides} slide(s) had no extractable text; they may be image-only or chart-heavy.")
    return ExtractedContent(kind="pptx", text="\n".join(parts), item_count=len(prs.slides), item_unit="slides", warnings=warnings)


def _row_values(row: Iterable[object]) -> list[str]:
    return [str(c).strip() for c in row if c is not None and str(c).strip()]


def extract_xlsx(data: bytes) -> ExtractedContent:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        non_empty_rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = _row_values(row)
            if cells:
                non_empty_rows.append(cells)
            if len(non_empty_rows) >= EXCEL_MAX_ROWS_PER_SHEET + 1:
                break
        parts.append(f"--- Sheet: {ws.title} ---")
        parts.append(f"Sheet summary: up to {ws.max_row or 0} rows x {ws.max_column or 0} columns; sampled {len(non_empty_rows)} non-empty row(s).")
        if not non_empty_rows:
            parts.append("(no non-empty cells found)")
            continue
        parts.append("Detected header / first non-empty row: " + " | ".join(non_empty_rows[0]))
        sample_rows = non_empty_rows[1:EXCEL_MAX_ROWS_PER_SHEET]
        if sample_rows:
            parts.append("Sample rows:")
            for idx, cells in enumerate(sample_rows, start=1):
                parts.append(f"Row {idx}: " + " | ".join(cells))
        if ws.max_row and ws.max_row > len(non_empty_rows):
            warnings.append(f"Sheet '{ws.title}' was sampled for prompt context; not every row was included.")
    return ExtractedContent(kind="xlsx", text="\n".join(parts), item_count=len(wb.worksheets), item_unit="sheets", warnings=warnings)


def extract_csv(data: bytes) -> ExtractedContent:
    text = data.decode("utf-8-sig", errors="replace")
    out: list[str] = []
    row_count = 0
    sampled = 0
    for row in _csv.reader(io.StringIO(text)):
        row_count += 1
        cells = [c.strip() for c in row if c and c.strip()]
        if cells and sampled < CSV_MAX_ROWS:
            sampled += 1
            prefix = "Header" if sampled == 1 else f"Row {sampled - 1}"
            out.append(prefix + ": " + " | ".join(cells))
    warnings = []
    if row_count > sampled:
        warnings.append(f"CSV was sampled: included {sampled} non-empty row(s) out of {row_count} total row(s).")
    return ExtractedContent(kind="csv", text="\n".join(out), item_count=row_count, item_unit="rows", warnings=warnings)


def extract_text(data: bytes) -> ExtractedContent:
    text = data.decode("utf-8-sig", errors="replace")
    return ExtractedContent(kind="text", text=text, item_count=len(text.splitlines()), item_unit="lines")


SUPPORTED_EXTENSIONS = ("pdf", "docx", "pptx", "xlsx", "csv", "txt", "md")
_EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "xlsx": extract_xlsx,
    "csv": extract_csv,
    "txt": extract_text,
    "md": extract_text,
}


def _weak_extraction_warnings(kind: str, raw_chars: int, item_count: int) -> list[str]:
    warnings: list[str] = []
    if raw_chars == 0:
        warnings.append("No extractable text found.")
    elif raw_chars < WEAK_EXTRACTION_CHAR_THRESHOLD:
        warnings.append(f"Weak extraction: only {raw_chars:,} character(s) found. The file may be short, scanned, image-heavy, or mostly charts.")
    if kind == "pptx" and item_count and raw_chars / max(item_count, 1) < 80:
        warnings.append("PowerPoint extraction is sparse for the number of slides; image/chart content may not be captured.")
    return warnings


def extract_file(filename: str, data: bytes) -> SourceDoc:
    ext = _ext(filename)
    file_size = len(data or b"")
    if ext not in _EXTRACTORS:
        supported = ", ".join(f".{x}" for x in SUPPORTED_EXTENSIONS)
        error = f"Unsupported file type '.{ext}' - skipped. Supported types: {supported}."
        return SourceDoc(filename=filename, kind="unknown", text="", note=error, extension=ext, file_size_bytes=file_size, errors=[error], weak_extraction=True)
    try:
        payload = _EXTRACTORS[ext](data)
    except Exception as exc:  # noqa: BLE001
        error = f"Could not parse: {shorten_text(exc, 240)}"
        kind = _EXTRACTORS[ext].__name__.replace("extract_", "") if callable(_EXTRACTORS[ext]) else ext
        return SourceDoc(filename=filename, kind=kind, text="", note=error, extension=ext, file_size_bytes=file_size, errors=[error], weak_extraction=True)

    raw = _clean_text(payload.text)
    raw_chars = len(raw)
    bounded, truncated, trunc_warning = _source_aware_truncate(raw)
    warnings = list(payload.warnings)
    warnings.extend(_weak_extraction_warnings(payload.kind, raw_chars, payload.item_count))
    if trunc_warning:
        warnings.append(trunc_warning)
    weak = raw_chars == 0 or raw_chars < WEAK_EXTRACTION_CHAR_THRESHOLD
    note_bits = []
    if truncated:
        note_bits.append("truncated")
    if warnings:
        note_bits.extend(warnings[:2])
        if len(warnings) > 2:
            note_bits.append(f"+{len(warnings) - 2} more warning(s)")
    if payload.errors:
        note_bits.extend(payload.errors[:1])
        if len(payload.errors) > 1:
            note_bits.append(f"+{len(payload.errors) - 1} more error(s)")

    return SourceDoc(
        filename=filename,
        kind=payload.kind,
        extension=ext,
        file_size_bytes=file_size,
        item_count=payload.item_count,
        item_unit=payload.item_unit,
        text=bounded,
        note="; ".join(note_bits),
        raw_char_count=raw_chars,
        extracted_char_count=len(bounded or ""),
        was_truncated=truncated,
        preview_text=(bounded or "")[:PREVIEW_CHAR_LIMIT],
        warnings=warnings,
        errors=payload.errors,
        weak_extraction=weak,
        ocr_attempted=payload.ocr_attempted,
        ocr_used=payload.ocr_used,
    )


def extract_many(files: List[tuple[str, bytes]]) -> List[SourceDoc]:
    return [extract_file(name, data) for name, data in files]


def docs_to_context(docs: List[SourceDoc]) -> str:
    """Concatenate parsed docs into a single labelled block for the prompt."""
    blocks = []
    for d in docs:
        header = f"### FILE: {d.filename} ({d.kind}; {d.item_count} {d.item_unit}; {d.extracted_character_count:,} chars sent"
        if d.was_truncated:
            header += "; source-aware truncated"
        if d.weak_extraction:
            header += "; weak extraction"
        header += ")"
        if d.warnings:
            header += "\nWARNINGS: " + " | ".join(d.warnings[:4])
        if d.errors:
            header += "\nERRORS: " + " | ".join(d.errors[:4])
        body = d.text.strip() if d.text.strip() else "(no extractable text)"
        blocks.append(f"{header}\n{body}")
    return "\n\n".join(blocks) if blocks else "(no files uploaded)"


def _message_summary(warnings: list[str], errors: list[str], *, limit: int = 180) -> str:
    parts: list[str] = []
    if warnings:
        parts.append(f"{len(warnings)} warning(s)")
    if errors:
        parts.append(f"{len(errors)} error(s)")
    messages = [*warnings, *errors]
    if messages:
        parts.append(shorten_text(messages[0], limit=limit))
        if len(messages) > 1:
            parts.append(f"+{len(messages) - 1} more")
    return " - ".join(parts) if parts else ""


def extraction_issue_rows(docs: List[SourceDoc]) -> list[dict]:
    """Detailed extraction warnings/errors for a collapsed UI expander/export."""
    rows: list[dict] = []
    for d in docs:
        for w in d.warnings or []:
            rows.append({"Filename": d.filename, "Severity": "Warning", "Message": sanitize_text(w)})
        for e in d.errors or []:
            rows.append({"Filename": d.filename, "Severity": "Error", "Message": sanitize_text(e)})
    return rows


def extraction_summary_rows(docs: List[SourceDoc]) -> list[dict]:
    rows: list[dict] = []
    for d in docs:
        warnings = [sanitize_text(w) for w in (d.warnings or [])]
        errors = [sanitize_text(e) for e in (d.errors or [])]
        processed = f"{d.item_count} {d.item_unit}"
        row = {
            "Filename": d.filename,
            "Type": d.kind,
            "Processed": processed,
            "Extracted chars": d.raw_character_count,
            "Prompt chars": d.extracted_character_count,
            "Truncated": "Yes" if d.was_truncated else "No",
            "Weak extraction": "Yes" if d.weak_extraction else "No",
            "Warning count": len(warnings),
            "Error count": len(errors),
            "Issue summary": _message_summary(warnings, errors),
            # Backward-compatible keys used by exports/tests.
            "Warnings/errors": _message_summary(warnings, errors),
            "filename": d.filename,
            "type": d.kind,
            "units": processed,
            "chars_extracted": d.raw_character_count,
            "chars_sent_to_model": d.extracted_character_count,
            "truncated": "Yes" if d.was_truncated else "No",
            "weak_extraction": "Yes" if d.weak_extraction else "No",
            "warning_count": len(warnings),
            "error_count": len(errors),
            "warnings": _message_summary(warnings, []),
            "errors": _message_summary([], errors),
        }
        rows.append(row)
    return rows


def extraction_report_markdown(docs: List[SourceDoc]) -> str:
    lines = ["# Phase 1 Extraction report", "", "Phase 1 extraction report", ""]
    for d in docs:
        lines.extend([
            f"## {d.filename}",
            f"- Type: {d.kind}",
            f"- Size: {d.file_size_bytes:,} bytes",
            f"- Processed: {d.item_count} {d.item_unit}",
            f"- Raw characters: {d.raw_character_count:,}",
            f"- Characters sent to model: {d.extracted_character_count:,}",
            f"- Truncated: {'Yes' if d.was_truncated else 'No'}",
            f"- Weak extraction: {'Yes' if d.weak_extraction else 'No'}",
            f"- OCR attempted: {'Yes' if d.ocr_attempted else 'No'}",
            f"- OCR used: {'Yes' if d.ocr_used else 'No'}",
        ])
        if d.warnings:
            lines.append("- Warnings:")
            lines.extend(f"  - {sanitize_text(w)}" for w in d.warnings)
        if d.errors:
            lines.append("- Errors:")
            lines.extend(f"  - {sanitize_text(e)}" for e in d.errors)
        lines.extend(["", "### Extracted text", "", sanitize_text(d.text) or "(no extractable text)", ""])
    return "\n".join(lines)


def extraction_report_xlsx_bytes(docs: List[SourceDoc]) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extraction_Summary"
    headers = ["Filename", "Type", "Processed", "Extracted chars", "Prompt chars", "Truncated", "Weak extraction", "Warnings/errors"]
    ws.append(excel_safe_row(headers))
    for row in extraction_summary_rows(docs):
        ws.append(excel_safe_row([row.get(h, "") for h in headers]))

    ws2 = wb.create_sheet("Extracted_Text")
    ws2.append(excel_safe_row(["Filename", "Extracted_Text_Sent_To_Model"]))
    for d in docs:
        ws2.append(excel_safe_row([d.filename, d.text]))

    ws3 = wb.create_sheet("Extraction_Warnings")
    ws3.append(excel_safe_row(["Filename", "Severity", "Message"]))
    for d in docs:
        for w in d.warnings:
            ws3.append(excel_safe_row([d.filename, "Warning", w]))
        for e in d.errors:
            ws3.append(excel_safe_row([d.filename, "Error", e]))
    if ws3.max_row == 1:
        ws3.append(excel_safe_row(["", "", "No extraction warnings or errors."]))

    fill = PatternFill("solid", fgColor="071D49")
    font = Font(color="F7F9FC", bold=True)
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = font
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 26
    ws.column_dimensions["H"].width = 70
    ws2.column_dimensions["B"].width = 120
    ws3.column_dimensions["C"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Backward-compatible UI aliases.
build_extraction_report_text = extraction_report_markdown
build_extraction_report_xlsx = extraction_report_xlsx_bytes
