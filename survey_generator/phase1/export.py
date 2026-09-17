"""Phase 1 export helpers."""
from __future__ import annotations

import io
from typing import Any, Iterable

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .models import SourceDoc
from .extract import extraction_summary_rows, extraction_report_markdown, extraction_report_xlsx_bytes
from survey_generator.runtime import excel_safe_row, sanitize_text


def _style_workbook(wb: openpyxl.Workbook) -> None:
    fill = PatternFill("solid", fgColor="071D49")
    font = Font(color="F7F9FC", bold=True)
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = font
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = 24


def _join(value: Any) -> str:
    if isinstance(value, list):
        return sanitize_text("; ".join(str(v).strip() for v in value if str(v).strip()))
    return sanitize_text(value or "")


def hypotheses_export_xlsx_bytes(
    hypotheses: Iterable[dict[str, Any]],
    docs: list[SourceDoc],
    validation_issues: Iterable[dict[str, Any]] | None = None,
) -> bytes:
    """Build a Phase 1 review/export workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hypotheses"
    headers = [
        "hypothesis_id", "review_status", "priority", "text", "business_question", "audience",
        "category", "rationale", "review_notes", "study_type", "source_files", "source_excerpt",
        "validation_warnings", "duplicate_of",
    ]
    ws.append(excel_safe_row(headers))
    for h in hypotheses:
        ws.append(excel_safe_row([_join(h.get(col, "")) for col in headers]))

    ws_files = wb.create_sheet("Source_Files")
    file_headers = ["Filename", "Type", "Processed", "Extracted chars", "Prompt chars", "Truncated", "Weak extraction", "Warnings/errors"]
    ws_files.append(excel_safe_row(file_headers))
    for r in extraction_summary_rows(docs):
        ws_files.append(excel_safe_row([r.get(h, "") for h in file_headers]))

    ws_warn = wb.create_sheet("Extraction_Warnings")
    ws_warn.append(excel_safe_row(["Filename", "Severity", "Message"]))
    for d in docs:
        for w in d.warnings:
            ws_warn.append(excel_safe_row([d.filename, "Warning", w]))
        for e in d.errors:
            ws_warn.append(excel_safe_row([d.filename, "Error", e]))
    if ws_warn.max_row == 1:
        ws_warn.append(excel_safe_row(["", "", "No extraction warnings or errors."]))

    ws_val = wb.create_sheet("Hypothesis_Validation")
    ws_val.append(excel_safe_row(["Row", "Hypothesis_ID", "Severity", "Field", "Message"]))
    for issue in validation_issues or []:
        ws_val.append(excel_safe_row([
            issue.get("row", ""), issue.get("hypothesis_id", ""), issue.get("severity", ""),
            issue.get("field", ""), issue.get("message", ""),
        ]))
    if ws_val.max_row == 1:
        ws_val.append(excel_safe_row(["", "", "", "", "No hypothesis validation issues."]))

    _style_workbook(wb)
    ws.column_dimensions["D"].width = 70
    ws.column_dimensions["H"].width = 55
    ws.column_dimensions["L"].width = 55
    ws_files.column_dimensions["H"].width = 70
    ws_warn.column_dimensions["C"].width = 80
    ws_val.column_dimensions["E"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Backward-compatible names imported by app.py and earlier tests.
def build_extraction_report_text(docs: Iterable[SourceDoc]) -> str:
    return extraction_report_markdown(list(docs))


def build_extraction_report_xlsx(docs: Iterable[SourceDoc]) -> bytes:
    return extraction_report_xlsx_bytes(list(docs))


def build_hypothesis_export_xlsx(
    hypotheses: Iterable[dict[str, Any]],
    docs: list[SourceDoc],
    validation_issues: Iterable[dict[str, Any]] | None = None,
) -> bytes:
    return hypotheses_export_xlsx_bytes(hypotheses, docs, validation_issues)
