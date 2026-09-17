from __future__ import annotations

import openpyxl

from survey_generator.phase2 import pipeline
from survey_generator.phase2.excel_writer import write_survey_spec
from survey_generator.phase2.models import Hypothesis, SectionProposal, StudySetup, SurveyQuestion


class FakeClient:
    max_json_retries = 0

    def __init__(self):
        self.responses = [
            [{"hypothesis_id": "H001", "section_id": "S1", "section_title": "Drivers", "construct": "importance", "question_type": "SINGLE_SELECT", "priority": "primary"}],
            [{"hypothesis_id": "H001", "section_id": "S1", "section_title": "Drivers", "construct": "importance", "question_type": "SINGLE_SELECT", "priority": "primary"}],
            [{"q_id": "TEMP1", "section": "Drivers", "section_id": "S1", "section_title": "Drivers", "hypothesis_ids": ["H001"], "question_text": "Which driver matters most?", "question_type": "SINGLE_SELECT", "options": ["[1 | Efficacy]", "[2 | Safety]"]}],
            [],
            [],
            [],
            {"revision_list": [], "escalated_issues": [], "auto_fix_issues": [], "major_issue_count": 0},
        ]

    def complete_json(self, _system, _user, expected_type=None):
        if not self.responses:
            raise AssertionError("No fake responses left")
        return self.responses.pop(0)

    def usage_snapshot(self):
        return {"calls": 7, "estimated_total_tokens": 100}


def test_mocked_full_phase2_pipeline_and_workbook_contract(tmp_path):
    study = StudySetup(study_type="ATU", target_population="HCPs", client_name="Client", study_title="Study")
    hyp = Hypothesis(hypothesis_id="H001", text="Adoption is higher when efficacy matters", category="Drivers")
    section = SectionProposal(section_id="S1", section_title="Drivers")
    result = pipeline.run_pipeline_after_gate(FakeClient(), study, [hyp], [section])
    assert result["questions"]
    assert result["validation"].status in {"PASS", "WARNING"}
    spec = tmp_path / "spec.xlsx"
    write_survey_spec(str(spec), study, result["questions"], [hyp], [section], coverage_map=result["validation"].coverage_map)
    wb = openpyxl.load_workbook(str(spec))
    assert "Survey_Outline" in wb.sheetnames
    assert wb["Survey_Outline"].max_row == 2  # header + one question row
