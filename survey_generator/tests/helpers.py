from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from docx import Document
from pptx import Presentation

from survey_generator.contract import OUTLINE_COLS, WORKBOOK_CONTRACT_VERSION


def minimal_pdf_bytes(text: str = "Hello PDF Extraction Test") -> bytes:
    stream = f"BT /F1 24 Tf 100 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n",
        b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
        b"5 0 obj << /Length " + str(len(stream)).encode("ascii") + b" >> stream\n" + stream + b"\nendstream endobj\n",
    ]
    body = b"%PDF-1.4\n"
    offsets = [0]
    for obj in objects:
        offsets.append(len(body))
        body += obj
    xref_offset = len(body)
    xref = [b"xref\n0 6\n", b"0000000000 65535 f \n"]
    for off in offsets[1:]:
        xref.append(f"{off:010d} 00000 n \n".encode("ascii"))
    trailer = f"trailer << /Root 1 0 R /Size 6 >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    return body + b"".join(xref) + trailer


def docx_bytes(text: str = "DOCX extraction text") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def pptx_bytes(text: str = "PPTX extraction text") -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = text
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Col A", "Col B"])
    ws.append(["Alpha", "Beta"])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def make_stage3_workbook(path: Path, rows: list[dict[str, str]], *, include_metadata: bool = True) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey_Metadata"
    if include_metadata:
        ws.append(["Field", "Value", "Notes"])
        ws.append(["Contract_Version", WORKBOOK_CONTRACT_VERSION, ""])
        ws.append(["Study_Title", "Regression Study", ""])
        ws.append(["Client_Name", "Client", ""])
        ws.append([])
        ws.append(["Segment_ID", "Audience_Segment", "Target_N", "Qualification_Notes"])
        ws.append(["SEG1", "Primary", "100", ""])
    else:
        ws.append(["Field", "Value", "Notes"])
    outline = wb.create_sheet("Survey_Outline")
    outline.append(OUTLINE_COLS)
    for row in rows:
        outline.append([row.get(c, "") for c in OUTLINE_COLS])
    wb.save(path)
    return path


def valid_outline_row(qid="Q001", qtype="SINGLE_SELECT", section_id="S1", order="001") -> dict[str, str]:
    options = "[1 | First option]\n[2 | Second option]"
    grid_cols = ""
    dropdown_options = ""
    if qtype.startswith("RATING_7"):
        options = "[1 | Statement one]"
    elif qtype == "RATING_DUAL":
        options = "[1 | Left statement vs Right statement]"
    elif qtype == "NUMERIC":
        options = "[1 | Numeric response]"
    elif qtype == "PERCENT_ALLOCATION":
        options = "[1 | Item A]\n[2 | Item B]"
    elif qtype == "RANKING":
        options = "[1 | Item A]\n[2 | Item B]"
    elif qtype == "AWARENESS_USAGE":
        options = "[1 | Product A]"
    elif qtype == "DROPDOWN":
        options = "[1 | Row A]"
        grid_cols = "[COL1 | Dropdown 1]"
        dropdown_options = "[1 | Choice A]\n[2 | Choice B]"
    elif qtype == "OPEN_END":
        options = ""
    elif qtype == "GRID":
        options = "[1 | Row A]"
        grid_cols = "[1 | Column A]\n[2 | Column B]"
    elif qtype == "TPP_DISPLAY":
        options = "[1 | Display content]"
    elif qtype == "INSTRUCTION":
        options = "[1 | Instruction text]"
    return {
        "Section_ID": section_id,
        "Section_Title": "Main Section",
        "Display_Order": order,
        "Question_ID": qid,
        "Question_Text": "Question text for " + qtype,
        "Question_Type": qtype,
        "Statements_or_Options": options,
        "Grid_Columns": grid_cols,
        "Dropdown_Options": dropdown_options,
        "Theme": "",
        "Audience_Split": "ALL",
        "Programming_Instructions": "",
        "Question_Skip_Logic": "",
        "Option_Level_Logic": "",
        "Termination_Logic": "",
        "Validation_Rule": "1|2|3|4|5|6|7" if qtype.startswith("RATING_7") else "",
        "Section_Objective": "Objective",
        "Notes": "",
        "Scale_Min": "1" if qtype.startswith("RATING_7") else "",
        "Scale_Max": "7" if qtype.startswith("RATING_7") else "",
        "Scale_Anchor_Low": "Low" if qtype.startswith("RATING_7") else "",
        "Scale_Anchor_High": "High" if qtype.startswith("RATING_7") else "",
        "Routing_Next": "",
        "Source_Hypotheses": "H001",
        "Interviewer_Notes": "",
    }
