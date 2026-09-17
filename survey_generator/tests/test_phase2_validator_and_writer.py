from __future__ import annotations

import openpyxl

from survey_generator.phase2.excel_writer import write_survey_spec
from survey_generator.phase2.models import HMapping, Hypothesis, RoutingSpec, ScaleSpec, SectionProposal, StudySetup, SurveyQuestion
from survey_generator.phase2.validator import validate
from survey_generator.contract import OUTLINE_COLS


def _base_objects():
    study = StudySetup(study_type="ATU", target_population="HCPs", client_name="Client", study_title="Study")
    hyp = Hypothesis(hypothesis_id="H001", text="Adoption is higher when efficacy perceptions are stronger", category="Drivers")
    section = SectionProposal(section_id="S1", section_title="Drivers")
    mapping = HMapping(hypothesis_id="H001", section_id="S1", section_title="Drivers", construct="importance", question_type="RATING_7_IMPORTANCE")
    return study, hyp, section, mapping


def test_phase2_validator_flags_ids_routing_and_missing_scale_anchors():
    _study, hyp, section, mapping = _base_objects()
    q1 = SurveyQuestion(
        q_id="Q1", section="Drivers", section_id="S1", section_title="Drivers",
        hypothesis_ids=["H001"], question_text="Rate importance", question_type="RATING_7_IMPORTANCE",
        options=["[1 | Efficacy]"], scale=ScaleSpec(min=1, max=7), routing=RoutingSpec(condition="if selected", next="UNKNOWN"),
    )
    q2 = q1.model_copy(deep=True)
    q2.q_id = "Q1"
    report = validate([q1, q2], [mapping], [hyp.hypothesis_id], [section])
    assert report.status == "BLOCKER"
    assert any("Duplicate q_id" in b or "duplicate" in b.lower() for b in report.blockers)
    assert any("unknown question" in b.lower() or "unknown" in b.lower() for b in report.blockers)
    assert any("scale anchors" in w.lower() for w in report.warnings)


def test_excel_writer_preserves_stage3_contract_fields(tmp_path):
    study, hyp, section, _mapping = _base_objects()
    q = SurveyQuestion(
        q_id="S1_Q01", section="Drivers", section_id="S1", section_title="Drivers",
        hypothesis_ids=["H001"], question_text="Rate importance", question_type="RATING_7_IMPORTANCE",
        options=["[1 | Efficacy]"], scale=ScaleSpec(min=1, max=7, anchor_low="Not important", anchor_high="Very important"),
        routing=RoutingSpec(condition="If not aware", next="SCREENOUT"), termination_logic=["Terminate if not aware"],
        interviewer_notes="Probe carefully",
    )
    path = tmp_path / "spec.xlsx"
    write_survey_spec(str(path), study, [q], [hyp], [section], coverage_map={"H001": ["S1_Q01"]})
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Survey_Outline"]
    headers = [c.value for c in ws[1]]
    row = {headers[i]: ws.cell(row=2, column=i+1).value for i in range(len(headers))}
    for col in ["Scale_Min", "Scale_Max", "Scale_Anchor_Low", "Scale_Anchor_High", "Routing_Next", "Termination_Logic", "Source_Hypotheses"]:
        assert col in OUTLINE_COLS
        assert row[col] not in (None, "")
    assert row["Routing_Next"] == "SCREENOUT"
    assert "H001" in row["Source_Hypotheses"]
