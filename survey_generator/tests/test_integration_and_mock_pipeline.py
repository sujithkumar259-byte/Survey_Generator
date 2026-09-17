from __future__ import annotations

import openpyxl

from survey_generator.phase2 import pipeline
from survey_generator.phase2.excel_writer import write_survey_spec
from survey_generator.phase2.models import Hypothesis, SectionProposal, StudySetup


class MockClient:
    max_json_retries = 0
    provider = "Mock"
    model = "mock-model"

    def __init__(self):
        self.calls = 0

    def usage_snapshot(self):
        return {"provider": self.provider, "model": self.model, "calls": self.calls}

    def complete_json(self, system, user, expected_type=None):
        self.calls += 1
        text = f"{system}\n{user}".lower()
        if "hypothesis mapper" in text or "map" in text and "section_id" in text:
            return [{
                "hypothesis_id": "H001",
                "section_id": "A",
                "section_title": "Drivers",
                "construct": "importance",
                "question_type": "RATING_7_IMPORTANCE",
                "priority": "primary",
            }]
        if "pressure" in text:
            return [{
                "hypothesis_id": "H001",
                "section_id": "A",
                "section_title": "Drivers",
                "construct": "importance",
                "question_type": "RATING_7_IMPORTANCE",
                "priority": "primary",
            }]
        if "draft" in text or "question drafter" in text:
            return [{
                "q_id": "TEMP1",
                "section": "Drivers",
                "section_id": "A",
                "section_title": "Drivers",
                "hypothesis_ids": ["H001"],
                "question_text": "How important are the following factors?",
                "question_type": "RATING_7_IMPORTANCE",
                "options": ["[1 | Efficacy]", "[2 | Safety]"],
                "scale": {"min": 1, "max": 7, "anchor_low": "Not important", "anchor_high": "Very important"},
                "routing": {"condition": "all respondents", "next": "END"},
            }]
        if "consolidator" in text or "triage" in text or "revision_list" in text:
            return {"revision_list": [], "escalated_issues": [], "auto_fix_issues": [], "major_issue_count": 0}
        # Domain expert / flow checker / quality reviewer issue lists.
        return []


def _study():
    return StudySetup(study_type="ATU", target_population="HCPs", client_name="Client", study_title="Study")


def _hypotheses():
    return [Hypothesis(hypothesis_id="H001", text="Physicians are more likely to adopt product X when efficacy is rated highly", category="Drivers")]


def _sections():
    return [SectionProposal(section_id="A", section_title="Drivers", rationale="", is_must_have=True)]


def test_mocked_full_phase2_pipeline_and_workbook_contract(tmp_path):
    client = MockClient()
    result = pipeline.run_pipeline_after_gate(client, _study(), _hypotheses(), _sections())
    assert result["questions"]
    assert result["validation"].status in {"PASS", "WARNING"}
    out = tmp_path / "generated.xlsx"
    write_survey_spec(
        str(out),
        _study(),
        result["questions"],
        _hypotheses(),
        _sections(),
        coverage_map=result["validation"].coverage_map,
        validation_report=result["validation"],
        final_approved=True,
    )
    wb = openpyxl.load_workbook(str(out))
    assert "Survey_Outline" in wb.sheetnames
    assert wb["Survey_Outline"].max_row == 2  # header + one question row
