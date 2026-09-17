from __future__ import annotations

from openpyxl import load_workbook

from survey_generator.phase2.excel_writer import write_survey_spec
from survey_generator.phase2.models import (
    HMapping, Hypothesis, RoutingSpec, ScaleSpec, SectionProposal, StudySetup, SurveyQuestion,
)
from survey_generator.phase2.validator import validate


def _study():
    return StudySetup(study_type="ATU", target_population="HCPs", client_name="Client", study_title="Study")


def _section():
    return SectionProposal(section_id="A", section_title="Drivers", rationale="", is_must_have=True)


def _hyp():
    return Hypothesis(hypothesis_id="H001", text="Access is associated with higher adoption", category="Drivers & barriers")


def _mapping():
    return HMapping(hypothesis_id="H001", section_id="A", section_title="Drivers", construct="attitude", question_type="RATING_7")


def test_phase2_validator_catches_blockers_and_routing():
    bad = SurveyQuestion(q_id="Q1", section="Drivers", section_id="A", question_text="", question_type="SINGLE_SELECT", options=["[1 | Only one]"])
    report = validate([bad], [_mapping()], ["H001"], [_section()])
    assert report.status == "BLOCKER"
    assert any("requires at least 2 options" in b for b in report.blockers)
    assert any("not covered" in b for b in report.blockers)

    q2 = SurveyQuestion(
        q_id="Q2", section="Drivers", section_id="A", question_text="Question?", question_type="SINGLE_SELECT",
        options=["[1 | Yes]", "[2 | No]"], hypothesis_ids=["H001"], routing=RoutingSpec(condition="If yes", next="NO_SUCH_Q"),
    )
    report2 = validate([q2], [_mapping()], ["H001"], [_section()])
    assert any("unknown target" in b for b in report2.blockers)


def test_excel_writer_preserves_scale_routing_termination_and_sources(tmp_path):
    q = SurveyQuestion(
        q_id="Q1", section="Drivers", section_id="A", section_title="Drivers", question_text="Please rate.",
        question_type="RATING_7", options=["[1 | Access is simple]"], hypothesis_ids=["H001"],
        scale=ScaleSpec(min=1, max=7, anchor_low="Not simple", anchor_high="Very simple"),
        routing=RoutingSpec(condition="If low", next="SCREENOUT"), termination_logic=["Low access terminates"],
    )
    path = tmp_path / "spec.xlsx"
    write_survey_spec(str(path), _study(), [q], [_hyp()], [_section()], coverage_map={"H001": ["Q1"]})
    ws = load_workbook(path, data_only=True)["Survey_Outline"]
    headers = [c.value for c in ws[1]]
    row = {headers[i]: ws.cell(row=2, column=i + 1).value for i in range(len(headers))}
    assert row["Scale_Anchor_Low"] == "Not simple"
    assert row["Scale_Anchor_High"] == "Very simple"
    assert row["Routing_Next"] == "SCREENOUT"
    assert "Low access terminates" in row["Termination_Logic"]
    assert row["Source_Hypotheses"] == "H001"
