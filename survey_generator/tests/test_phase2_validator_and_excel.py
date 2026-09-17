from __future__ import annotations

import openpyxl

from survey_generator.phase2.excel_writer import write_survey_spec
from survey_generator.phase2.models import (
    HMapping,
    Hypothesis,
    RoutingSpec,
    ScaleSpec,
    SectionProposal,
    StudySetup,
    SurveyQuestion,
)
from survey_generator.phase2.validator import validate


def _study():
    return StudySetup(study_type="ATU", target_population="HCPs", client_name="Client", study_title="Study")


def _hyp():
    return Hypothesis(hypothesis_id="H001", text="Physicians are more likely to adopt product X when efficacy is rated highly", category="Drivers")


def _section():
    return SectionProposal(section_id="A", section_title="Drivers", rationale="", is_must_have=True)


def _mapping():
    return HMapping(hypothesis_id="H001", section_id="A", section_title="Drivers", construct="importance", question_type="RATING_7_IMPORTANCE")


def _question(**overrides):
    data = dict(
        q_id="A_Q01",
        section="Drivers",
        section_id="A",
        section_title="Drivers",
        hypothesis_ids=["H001"],
        question_text="How important are the following factors?",
        question_type="RATING_7_IMPORTANCE",
        options=["[1 | Efficacy]", "[2 | Safety]"],
        scale=ScaleSpec(min=1, max=7, anchor_low="Not important", anchor_high="Very important"),
        routing=RoutingSpec(condition="if A", next="END"),
        termination_logic=["Terminate if not qualified"],
        interviewer_notes="Probe if needed",
    )
    data.update(overrides)
    return SurveyQuestion(**data)


def test_phase2_validator_catches_core_blockers():
    q1 = _question(q_id="A_Q01")
    q2 = _question(q_id="A_Q01", routing=RoutingSpec(condition="if selected", next="UNKNOWN"), question_type="SINGLE_SELECT", options=["[1 | Only one]"])
    report = validate([q1, q2], [_mapping()], ["H001", "H999"], [_section()])
    text = " ".join(report.blockers)
    assert report.status == "BLOCKER"
    assert "not covered" in text
    assert "appears" in text
    assert "unknown target" in text
    assert "requires at least 2 options" in text


def test_excel_writer_preserves_production_fields(tmp_path):
    out = tmp_path / "spec.xlsx"
    q = _question()
    report = validate([q], [_mapping()], ["H001"], [_section()])
    write_survey_spec(str(out), _study(), [q], [_hyp()], [_section()], coverage_map=report.coverage_map, validation_report=report, final_approved=True)
    wb = openpyxl.load_workbook(out, data_only=True)
    assert "Survey_Metadata" in wb.sheetnames
    assert "Survey_Outline" in wb.sheetnames
    ws = wb["Survey_Outline"]
    headers = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
    assert row["Scale_Anchor_Low"] == "Not important"
    assert row["Scale_Anchor_High"] == "Very important"
    assert row["Routing_Next"] == "END"
    assert "Terminate" in row["Termination_Logic"]
    assert row["Source_Hypotheses"] == "H001"
    assert ws.max_row == 2  # header + one question row
